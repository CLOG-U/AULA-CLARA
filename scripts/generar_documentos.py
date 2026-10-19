#!/usr/bin/env python3
"""Create the stage-2 Word and a combined Word preserving every original stage-1 page.

Requires python-docx and PyMuPDF. Screenshots must come from the running Android app.
Usage: python scripts/generar_documentos.py --etapa1 path/to/etapa1.pdf --salida path/to/output
"""
import argparse
from pathlib import Path

import fitz
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]


def add_link(doc, text, url):
    paragraph = doc.add_paragraph()
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), doc.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    properties = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '007F78')
    underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single')
    properties.append(color); properties.append(underline); run.append(properties)
    label = OxmlElement('w:t'); label.text = text; run.append(label)
    link.append(run); paragraph._p.append(link)


def configure(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.8)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.text = 'Aula Clara · Pantallas y navegación · Etapa 2'


def add_stage2(doc, images):
    """Include only the student details, repository URL and three requested captures."""
    doc.core_properties.author = 'Carlos Luis Ortiz García'
    doc.core_properties.title = 'Aula Clara · Trabajo autónomo'
    doc.core_properties.subject = 'Pantallas y navegación con Jetpack Compose'
    configure(doc.sections[-1])
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    for name in ['Title', 'Heading 1', 'Heading 2']:
        doc.styles[name].font.color.rgb = RGBColor.from_string('007F78')
    doc.add_heading('Aula Clara', 0)
    doc.add_paragraph('Segunda etapa · Pantallas y navegación', style='Subtitle')
    doc.add_paragraph('Integrante único: Carlos Luis Ortiz García')
    doc.add_heading('Repositorio del proyecto', 1)
    add_link(doc, 'https://github.com/CLOG-U/AULA-CLARA', 'https://github.com/CLOG-U/AULA-CLARA')
    heading = doc.add_heading('Capturas de pantalla del MVP', 1)
    heading.paragraph_format.keep_with_next = True
    captions = [('01-inicio.png', 'Figura 1. Inicio de Aula Clara.'),
                ('02-formulario.png', 'Figura 2. Formulario de entrega.'),
                ('03-detalle-completado.png', 'Figura 3. Detalle de una entrega completada.')]
    for index, (name, caption) in enumerate(captions):
        image = images / name
        if not image.is_file():
            raise FileNotFoundError('Missing real Android screenshot: ' + str(image))
        doc.add_picture(str(image), height=Inches(5.8 if index == 0 else 7.0))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.paragraphs[-1].paragraph_format.page_break_before = index > 0
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
        doc.add_paragraph(caption)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--etapa1', type=Path, required=True)
    parser.add_argument('--salida', type=Path, required=True)
    args = parser.parse_args()
    args.salida.mkdir(parents=True, exist_ok=True)
    evidence = ROOT / 'docs/evidencias'
    doc = Document()
    add_stage2(doc, evidence)
    doc.save(args.salida / 'Aula_Clara_Etapa_2.docx')
    combined = Document()
    section = combined.sections[0]
    source = fitz.open(args.etapa1)
    first = source[0].rect
    section.page_width, section.page_height = Pt(first.width), Pt(first.height)
    section.top_margin = section.bottom_margin = Pt(1)
    section.left_margin = section.right_margin = Pt(1)
    cache = args.salida / '.paginas-etapa1'
    cache.mkdir(exist_ok=True)
    for i, page in enumerate(source):
        path = cache / f'{i+1:02d}.png'
        page.get_pixmap(matrix=fitz.Matrix(1.6, 1.6)).save(path)
        paragraph = combined.add_paragraph()
        paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1
        paragraph.paragraph_format.page_break_before = i > 0
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.font.size = Pt(1)
        run.add_picture(str(path), height=Pt(first.height - 26))
    combined.add_section(WD_SECTION_START.NEW_PAGE)
    add_stage2(combined, evidence)
    combined.save(args.salida / 'Aula_Clara_Etapas_1_y_2.docx')
    (args.salida / 'LEER_ENTREGA.txt').write_text(
        'Aula_Clara_Etapas_1_y_2.docx reúne las 32 páginas originales de la etapa 1 y la etapa 2.\n'
        'Como la etapa 1 se proporcionó en PDF, sus páginas se conservaron como imágenes fieles; la etapa 2 es texto editable.\n'
        'Aula_Clara_Etapa_2.docx contiene únicamente la sección nueva editable para integrarla en el Word original si se dispone de él.\n'
        'Integrante único declarado: Carlos Luis Ortiz García.\n'
        'Revisar el estado de publicación del repositorio y subir personalmente el Word a Moodle.\n', encoding='utf-8')
    print('Generated Word documents in', args.salida)


if __name__ == '__main__':
    main()
